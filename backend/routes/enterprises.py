from flask import request, jsonify, send_file
from extensions import get_db, record_operation_log, token_required, user_can
from datetime import datetime, timedelta
from io import BytesIO
import json, re

from . import enterprises_bp


@enterprises_bp.route('/api/enterprises', methods=['GET'])
@token_required
def get_enterprises():
    """获取企业信息库列表。"""
    payload = request.current_user
    role = payload['role']

    keyword = request.args.get('keyword', '')
    relationship_status = request.args.get('relationship_status', '')
    owner_id = request.args.get('owner_id', '')

    db = get_db()
    cursor = db.cursor()

    conditions = []
    params = []

    if not user_can(payload['username'], 'data.view_all'):
        conditions.append("e.owner_id = ?")
        params.append(payload['username'])

    if keyword:
        conditions.append("(e.name LIKE ? OR e.contact_person LIKE ? OR e.brief LIKE ?)")
        params.extend([f'%{keyword}%', f'%{keyword}%', f'%{keyword}%'])
    if relationship_status:
        conditions.append("e.relationship_status = ?")
        params.append(relationship_status)
    if owner_id:
        conditions.append("e.owner_id = ?")
        params.append(owner_id)

    where_clause = ' AND '.join(conditions) if conditions else '1=1'

    cursor.execute(f"""
        SELECT e.*, u.name as owner_name,
               (SELECT COUNT(*) FROM enterprise_visits ev WHERE ev.enterprise_id = e.id) as visit_count,
               (SELECT COUNT(*) FROM customers c WHERE c.company LIKE '%' || e.name || '%' OR c.company = e.name) as customer_count,
               (SELECT COUNT(*) FROM business b
                WHERE b.cust_id IN (SELECT c.id FROM customers c WHERE c.company LIKE '%' || e.name || '%' OR c.company = e.name)) as business_count,
               (SELECT COUNT(*) FROM contracts ct
                WHERE ct.cust_id IN (SELECT c.id FROM customers c WHERE c.company LIKE '%' || e.name || '%' OR c.company = e.name)) as contract_count
        FROM enterprises e
        LEFT JOIN users u ON e.owner_id = u.username
        WHERE {where_clause}
        ORDER BY e.updated_at DESC
    """, params)

    rows = cursor.fetchall()
    enterprises = [dict(row) for row in rows]

    return jsonify({'code': 200, 'message': 'success', 'data': enterprises})


@enterprises_bp.route('/api/enterprises/<int:enterprise_id>', methods=['GET'])
@token_required
def get_enterprise(enterprise_id):
    """获取单个企业详情，自动关联客户、商机、合同、拜访数据，打通数据链条。"""
    db = get_db()
    cursor = db.cursor()

    cursor.execute("""
        SELECT e.*, u.name as owner_name
        FROM enterprises e
        LEFT JOIN users u ON e.owner_id = u.username
        WHERE e.id = ?
    """, (enterprise_id,))
    row = cursor.fetchone()
    if not row:
        return jsonify({'code': 404, 'message': '企业不存在', 'data': None})

    enterprise = dict(row)
    ent_name = enterprise.get('name', '')

    # --- 1. 自动匹配客户：按企业名称匹配 customers.company ---
    matched_cust_ids = []
    if ent_name:
        cursor.execute("""
            SELECT id, name, company, phone, level, owner_id, last_follow, created_at
            FROM customers
            WHERE company LIKE ? OR company = ?
            ORDER BY last_follow DESC
        """, (f'%{ent_name}%', ent_name))
        customers = [dict(r) for r in cursor.fetchall()]
        enterprise['customers'] = customers
        matched_cust_ids = [c['id'] for c in customers]
    else:
        enterprise['customers'] = []

    # --- 2. 关联商机：通过匹配到的客户的 cust_id 查询 business ---
    if matched_cust_ids:
        placeholders = ','.join(['?'] * len(matched_cust_ids))
        cursor.execute(f"""
            SELECT b.*, c.company as customer_company, c.name as customer_name
            FROM business b
            LEFT JOIN customers c ON b.cust_id = c.id
            WHERE b.cust_id IN ({placeholders})
            ORDER BY b.created_at DESC
        """, matched_cust_ids)
        enterprise['business'] = [dict(r) for r in cursor.fetchall()]
    else:
        enterprise['business'] = []

    # --- 3. 关联合同：通过匹配到的客户的 cust_id 查询 contracts ---
    if matched_cust_ids:
        placeholders = ','.join(['?'] * len(matched_cust_ids))
        cursor.execute(f"""
            SELECT ct.*, c.company as customer_company
            FROM contracts ct
            LEFT JOIN customers c ON ct.cust_id = c.id
            WHERE ct.cust_id IN ({placeholders})
            ORDER BY ct.sign_date DESC
        """, matched_cust_ids)
        enterprise['contracts'] = [dict(r) for r in cursor.fetchall()]
    else:
        enterprise['contracts'] = []

    # --- 4. 关联拜访记录：手动关联(enterprise_visits) + 匹配客户的拜访记录 ---
    visit_ids_set = set()

    # 4a. enterprise_visits 关联表中的拜访
    cursor.execute("""
        SELECT v.*, u.name as visitor_name,
               c.company as customer_company
        FROM enterprise_visits ev
        JOIN visits v ON ev.visit_id = v.id
        LEFT JOIN users u ON v.visitor_id = u.username
        LEFT JOIN customers c ON v.cust_id = c.id
        WHERE ev.enterprise_id = ?
        ORDER BY v.plan_date DESC, v.plan_time DESC
    """, (enterprise_id,))
    manual_visits = [dict(r) for r in cursor.fetchall()]
    visit_ids_set.update(v['id'] for v in manual_visits)

    # 4b. 匹配客户的拜访记录（去重）
    auto_visits = []
    if matched_cust_ids:
        placeholders = ','.join(['?'] * len(matched_cust_ids))
        cursor.execute(f"""
            SELECT v.*, u.name as visitor_name,
                   c.company as customer_company
            FROM visits v
            LEFT JOIN users u ON v.visitor_id = u.username
            LEFT JOIN customers c ON v.cust_id = c.id
            WHERE v.cust_id IN ({placeholders})
            ORDER BY v.plan_date DESC, v.plan_time DESC
        """, matched_cust_ids)
        for r in cursor.fetchall():
            v = dict(r)
            if v['id'] not in visit_ids_set:
                auto_visits.append(v)
                visit_ids_set.add(v['id'])

    # 合并：手动关联标记 link_type='manual'，自动匹配标记 link_type='auto'
    for v in manual_visits:
        v['link_type'] = 'manual'
    for v in auto_visits:
        v['link_type'] = 'auto'
    enterprise['visits'] = manual_visits + auto_visits

    # 统计摘要
    enterprise['summary'] = {
        'customer_count': len(enterprise['customers']),
        'business_count': len(enterprise['business']),
        'contract_count': len(enterprise['contracts']),
        'visit_count': len(enterprise['visits']),
        'business_total_amount': sum(b.get('amount', 0) or 0 for b in enterprise['business']),
        'contract_total_amount': sum(c.get('total_amt', 0) or 0 for c in enterprise['contracts']),
        'contract_paid_amount': sum(c.get('paid_amt', 0) or 0 for c in enterprise['contracts']),
    }

    return jsonify({'code': 200, 'message': 'success', 'data': enterprise})


@enterprises_bp.route('/api/enterprises', methods=['POST'])
@token_required
def create_enterprise():
    """新建企业信息。"""
    payload = request.current_user
    data = request.get_json(silent=True) or {}

    if not data.get('name'):
        return jsonify({'code': 400, 'message': '企业名称不能为空', 'data': None})

    db = get_db()
    cursor = db.cursor()

    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    cursor.execute("""
        INSERT INTO enterprises (name, established_date, location, personnel_size, brief,
            registered_capital, business_scope, main_qualifications, main_products,
            relationship_status, cooperation_opportunities, website, contact_person,
            contact_info, owner_id, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        data.get('name', ''),
        data.get('established_date', ''),
        data.get('location', ''),
        data.get('personnel_size', ''),
        data.get('brief', ''),
        data.get('registered_capital', ''),
        data.get('business_scope', ''),
        data.get('main_qualifications', ''),
        data.get('main_products', ''),
        data.get('relationship_status', '未接触'),
        data.get('cooperation_opportunities', ''),
        data.get('website', ''),
        data.get('contact_person', ''),
        data.get('contact_info', ''),
        data.get('owner_id', '') or payload['username'],
        now, now,
    ))
    db.commit()
    enterprise_id = cursor.lastrowid

    record_operation_log(payload['username'], '新建', '企业信息库', f'新建企业: {data["name"]}')

    return jsonify({'code': 200, 'message': '创建成功', 'data': {'id': enterprise_id}})


@enterprises_bp.route('/api/enterprises/<int:enterprise_id>', methods=['PUT'])
@token_required
def update_enterprise(enterprise_id):
    """修改企业信息。"""
    payload = request.current_user
    data = request.get_json(silent=True) or {}

    db = get_db()
    cursor = db.cursor()

    cursor.execute("SELECT id FROM enterprises WHERE id = ?", (enterprise_id,))
    if not cursor.fetchone():
        return jsonify({'code': 404, 'message': '企业不存在', 'data': None})

    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    fields = ['name', 'established_date', 'location', 'personnel_size', 'brief',
              'registered_capital', 'business_scope', 'main_qualifications', 'main_products',
              'relationship_status', 'cooperation_opportunities', 'website', 'contact_person',
              'contact_info', 'owner_id']

    sets = []
    params = []
    for f in fields:
        if f in data:
            sets.append(f"{f} = ?")
            params.append(data[f])
    sets.append("updated_at = ?")
    params.append(now)
    params.append(enterprise_id)

    cursor.execute(f"UPDATE enterprises SET {', '.join(sets)} WHERE id = ?", params)
    db.commit()

    record_operation_log(payload['username'], '修改', '企业信息库', f'修改企业ID: {enterprise_id}')

    return jsonify({'code': 200, 'message': '修改成功', 'data': None})


@enterprises_bp.route('/api/enterprises/<int:enterprise_id>', methods=['DELETE'])
@token_required
def delete_enterprise(enterprise_id):
    """删除企业信息。"""
    payload = request.current_user

    db = get_db()
    cursor = db.cursor()

    cursor.execute("SELECT name FROM enterprises WHERE id = ?", (enterprise_id,))
    row = cursor.fetchone()
    if not row:
        return jsonify({'code': 404, 'message': '企业不存在', 'data': None})

    name = row['name']
    cursor.execute("DELETE FROM enterprise_visits WHERE enterprise_id = ?", (enterprise_id,))
    cursor.execute("DELETE FROM enterprises WHERE id = ?", (enterprise_id,))
    db.commit()

    record_operation_log(payload['username'], '删除', '企业信息库', f'删除企业: {name}')

    return jsonify({'code': 200, 'message': '删除成功', 'data': None})


@enterprises_bp.route('/api/enterprises/<int:enterprise_id>/visits', methods=['POST'])
@token_required
def link_visit(enterprise_id):
    """关联拜访记录到企业。"""
    data = request.get_json(silent=True) or {}
    visit_id = data.get('visit_id')
    if not visit_id:
        return jsonify({'code': 400, 'message': '请选择拜访记录', 'data': None})

    db = get_db()
    cursor = db.cursor()

    cursor.execute("SELECT id FROM enterprises WHERE id = ?", (enterprise_id,))
    if not cursor.fetchone():
        return jsonify({'code': 404, 'message': '企业不存在', 'data': None})

    cursor.execute("SELECT id FROM visits WHERE id = ?", (visit_id,))
    if not cursor.fetchone():
        return jsonify({'code': 404, 'message': '拜访记录不存在', 'data': None})

    try:
        cursor.execute(
            "INSERT INTO enterprise_visits (enterprise_id, visit_id) VALUES (?, ?)",
            (enterprise_id, visit_id)
        )
        db.commit()
    except Exception:
        return jsonify({'code': 409, 'message': '该拜访记录已关联', 'data': None})

    return jsonify({'code': 200, 'message': '关联成功', 'data': None})


@enterprises_bp.route('/api/enterprises/<int:enterprise_id>/visits/<int:visit_id>', methods=['DELETE'])
@token_required
def unlink_visit(enterprise_id, visit_id):
    """取消关联拜访记录。"""
    db = get_db()
    cursor = db.cursor()

    cursor.execute(
        "DELETE FROM enterprise_visits WHERE enterprise_id = ? AND visit_id = ?",
        (enterprise_id, visit_id)
    )
    db.commit()

    return jsonify({'code': 200, 'message': '取消关联成功', 'data': None})


@enterprises_bp.route('/api/enterprises/import', methods=['POST'])
@token_required
def import_enterprises():
    """批量导入企业信息。"""
    payload = request.current_user
    data = request.get_json(silent=True) or {}
    rows = data.get('rows', [])

    if not rows:
        return jsonify({'code': 400, 'message': '没有可导入的数据', 'data': None})

    db = get_db()
    cursor = db.cursor()
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

    success_count = 0
    fail_count = 0
    results = []

    for idx, row in enumerate(rows, 1):
        name = (row.get('name') or '').strip()
        if not name:
            fail_count += 1
            results.append({'row_index': idx, 'success': False, 'message': '企业名称不能为空'})
            continue

        try:
            cursor.execute("""
                INSERT INTO enterprises (name, established_date, location, personnel_size, brief,
                    registered_capital, business_scope, main_qualifications, main_products,
                    relationship_status, cooperation_opportunities, website, contact_person,
                    contact_info, owner_id, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                name,
                row.get('established_date', ''),
                row.get('location', ''),
                row.get('personnel_size', ''),
                row.get('brief', ''),
                row.get('registered_capital', ''),
                row.get('business_scope', ''),
                row.get('main_qualifications', ''),
                row.get('main_products', ''),
                row.get('relationship_status', '未接触'),
                row.get('cooperation_opportunities', ''),
                row.get('website', ''),
                row.get('contact_person', ''),
                row.get('contact_info', ''),
                row.get('owner_id', '') or payload['username'],
                now, now,
            ))
            success_count += 1
            results.append({'row_index': idx, 'success': True, 'message': '成功'})
        except Exception as e:
            fail_count += 1
            results.append({'row_index': idx, 'success': False, 'message': str(e)})

    db.commit()

    record_operation_log(payload['username'], '导入', '企业信息库',
                         f'导入企业信息: 成功{success_count}条, 失败{fail_count}条')

    return jsonify({
        'code': 200,
        'message': f'导入完成: 成功{success_count}条, 失败{fail_count}条',
        'data': {
            'success_count': success_count,
            'fail_count': fail_count,
            'results': results,
        }
    })


def register_routes(app):
    app.register_blueprint(enterprises_bp)


# ============================================================
# 1. 从拜访纪要同步企业到企业信息库
# ============================================================
@enterprises_bp.route('/api/enterprises/sync-from-knowledge', methods=['POST'])
@token_required
def sync_from_knowledge():
    """扫描知识库拜访纪要，提取企业信息同步到企业信息库。

    策略：
    - 有 cust_id 的：直接从 customers 表获取企业基本信息
    - 无 cust_id 的：用正则从 content 提取拜访单位名称
    - 按企业名称去重，已存在则跳过
    """
    payload = request.current_user
    db = get_db()
    cursor = db.cursor()

    # 收集所有拜访纪要中的企业信息
    enterprises_map = {}  # name -> {fields}

    # --- 1. knowledge_documents（doc_type='visit_summary'） ---
    cursor.execute("""
        SELECT kd.id, kd.title, kd.content, kd.cust_id, kd.owner_id, kd.created_at
        FROM knowledge_documents kd
        WHERE kd.doc_type = 'visit_summary'
    """)
    for r in cursor.fetchall():
        info = _extract_enterprise_from_doc(r, cursor)
        if info and info['name']:
            name = info['name']
            if name not in enterprises_map:
                enterprises_map[name] = info
            else:
                # 合并：保留非空字段
                for k, v in info.items():
                    if v and not enterprises_map[name].get(k):
                        enterprises_map[name][k] = v

    # --- 2. knowledge_base（category='visit_summary'） ---
    cursor.execute("""
        SELECT kb.id, kb.title, kb.content, kb.cust_id, kb.owner_id, kb.created_at
        FROM knowledge_base kb
        WHERE kb.category = 'visit_summary'
    """)
    for r in cursor.fetchall():
        info = _extract_enterprise_from_doc(r, cursor)
        if info and info['name']:
            name = info['name']
            if name not in enterprises_map:
                enterprises_map[name] = info
            else:
                for k, v in info.items():
                    if v and not enterprises_map[name].get(k):
                        enterprises_map[name][k] = v

    # --- 3. 同步到 enterprises 表（按名称去重） ---
    existing_names = {
        r['name'] for r in cursor.execute("SELECT name FROM enterprises").fetchall()
    }
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    added = 0
    skipped = 0
    for name, info in enterprises_map.items():
        if name in existing_names:
            skipped += 1
            continue
        cursor.execute("""
            INSERT INTO enterprises (name, location, contact_person, contact_info,
                business_scope, brief, relationship_status, owner_id,
                created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            name,
            info.get('location', ''),
            info.get('contact_person', ''),
            info.get('contact_info', ''),
            info.get('business_scope', ''),
            info.get('brief', ''),
            '已接触',  # 拜访纪要说明已接触
            info.get('owner_id', '') or payload['username'],
            now, now,
        ))
        added += 1

    db.commit()
    record_operation_log(
        payload['username'], '同步', '企业信息库',
        f'从拜访纪要同步企业: 新增{added}条, 跳过已存在{skipped}条'
    )
    return jsonify({
        'code': 200,
        'message': f'同步完成: 新增{added}条, 跳过已存在{skipped}条',
        'data': {'added': added, 'skipped': skipped, 'total_scanned': len(enterprises_map)},
    })


def _extract_enterprise_from_doc(row, cursor):
    """从拜访纪要文档中提取企业信息。"""
    name = ''
    contact_person = ''
    location = ''
    contact_info = ''
    brief = ''
    owner_id = row['owner_id'] or ''

    # 有 cust_id 的直接从 customers 表获取
    if row['cust_id']:
        cust = cursor.execute(
            "SELECT company, contact_name, phone, email, address, industry, region "
            "FROM customers WHERE id=?", (row['cust_id'],)
        ).fetchone()
        if cust:
            name = (cust['company'] or '').strip()
            contact_person = (cust['contact_name'] or '').strip()
            contact_info = '; '.join(filter(None, [cust['phone'], cust['email']]))
            location = cust['address'] or ''
            if name:
                return {
                    'name': name,
                    'contact_person': contact_person,
                    'contact_info': contact_info,
                    'location': location,
                    'business_scope': cust['industry'] or '',
                    'owner_id': owner_id,
                }

    # 无 cust_id 的用正则从 content 提取
    content = row['content'] or ''
    title = row['title'] or ''

    # 尝试从 title 提取（格式如 "20220809能建集团二公司交流纪要"）
    m = re.match(r'^\d{6,8}(.+)', title)
    if m and not name:
        name = m.group(1).replace('交流纪要', '').replace('拜访纪要', '').replace('拜访记录', '').strip()

    # 从 content 提取"拜访单位"字段
    if not name:
        m = re.search(r'拜访\s*单位[｜|\s]*([^\n｜|]+)', content)
        if m:
            name = m.group(1).strip().split('\n')[0].strip()

    # 提取对方参会人员作为联系人
    m = re.search(r'对方参会人员[｜|\s]*([^\n]+(?:\n[^\n｜|]+)*)', content)
    if m and not contact_person:
        contact_person = m.group(1).strip().replace('\n', '; ')[:100]

    # 提取地点
    m = re.search(r'地点[｜|\s]*([^\n｜|]+)', content)
    if m and not location:
        location = m.group(1).strip().split('\n')[0].strip()[:80]

    if not name:
        return None

    # 截断过长字段
    name = name[:100]
    return {
        'name': name,
        'contact_person': contact_person,
        'contact_info': contact_info,
        'location': location,
        'owner_id': owner_id,
    }


# ============================================================
# 2. 企业信息库导出 Excel
# ============================================================
@enterprises_bp.route('/api/enterprises/export', methods=['GET'])
@token_required
def export_enterprises():
    """导出企业信息库为 Excel。"""
    payload = request.current_user
    db = get_db()
    cursor = db.cursor()

    keyword = request.args.get('keyword', '')
    relationship_status = request.args.get('relationship_status', '')

    conditions = []
    params = []
    if not user_can(payload['username'], 'data.view_all'):
        conditions.append("e.owner_id = ?")
        params.append(payload['username'])
    if keyword:
        conditions.append("(e.name LIKE ? OR e.contact_person LIKE ? OR e.brief LIKE ?)")
        params.extend([f'%{keyword}%', f'%{keyword}%', f'%{keyword}%'])
    if relationship_status:
        conditions.append("e.relationship_status = ?")
        params.append(relationship_status)

    where_clause = ' AND '.join(conditions) if conditions else '1=1'

    cursor.execute(f"""
        SELECT e.name, e.established_date, e.location, e.personnel_size,
               e.registered_capital, e.relationship_status, e.contact_person,
               e.contact_info, e.website, e.brief, e.business_scope,
               e.main_qualifications, e.main_products, e.cooperation_opportunities,
               u.name as owner_name
        FROM enterprises e
        LEFT JOIN users u ON e.owner_id = u.username
        WHERE {where_clause}
        ORDER BY e.updated_at DESC
    """, params)
    rows = cursor.fetchall()

    try:
        from openpyxl import Workbook
    except ImportError:
        # 降级为 CSV
        import csv
        from io import StringIO
        headers = ['企业名称', '成立时间', '公司位置', '人员规模', '注册资本',
                   '关系状态', '联系人', '联系方式', '单位网址', '单位简介',
                   '业务范围', '主要资质', '主要产品', '合作机会点', '负责人']
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        for r in rows:
            writer.writerow([r[i] or '' for i in range(len(headers))])
        csv_bytes = output.getvalue().encode('utf-8-sig')
        return send_file(BytesIO(csv_bytes), mimetype='text/csv',
                         as_attachment=True, download_name='企业信息库.csv')

    wb = Workbook()
    ws = wb.active
    ws.title = '企业信息库'
    headers = ['企业名称', '成立时间', '公司位置', '人员规模', '注册资本',
               '关系状态', '联系人', '联系方式', '单位网址', '单位简介',
               '业务范围', '主要资质', '主要产品', '合作机会点', '负责人']
    ws.append(headers)
    for r in rows:
        ws.append([r[i] or '' for i in range(len(headers))])

    # 列宽
    widths = [25, 12, 18, 10, 12, 10, 10, 18, 18, 30, 25, 25, 25, 25, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True,
                     download_name=f'企业信息库_{datetime.now().strftime("%Y%m%d")}.xlsx')


# ============================================================
# 3. 历史客户拜访推荐算法
# ============================================================
@enterprises_bp.route('/api/enterprises/visit-recommendations', methods=['GET'])
@token_required
def visit_recommendations():
    """根据历史数据计算企业拜访推荐分数。

    评分维度（百分制）：
    - 商机价值（35%）：未成交商机金额 + 进行中商机数量
    - 合同价值（20%）：已有合同金额总和（高价值客户优先维护）
    - 拜访间隔（25%）：距上次拜访越久分越高（90天以上满分）
    - 关系状态（10%）：已接触/合作中 > 未接触
    - 拜访频率（10%）：历史拜访次数适中分高

    返回：按分数降序的推荐列表，含推荐理由。
    """
    payload = request.current_user
    db = get_db()
    cursor = db.cursor()

    # 获取当前用户负责的企业（或全部，取决于权限）
    can_all = user_can(payload['username'], 'data.view_all')

    base_sql = """
        SELECT e.id, e.name, e.relationship_status, e.contact_person, e.contact_info,
               e.location, e.owner_id, u.name as owner_name,
               (SELECT MAX(v.plan_date) FROM visits v
                  LEFT JOIN enterprise_visits ev ON ev.visit_id = v.id
                  WHERE ev.enterprise_id = e.id) as last_visit_date,
               (SELECT COUNT(*) FROM visits v
                  LEFT JOIN enterprise_visits ev ON ev.visit_id = v.id
                  WHERE ev.enterprise_id = e.id) as visit_count
        FROM enterprises e
        LEFT JOIN users u ON e.owner_id = u.username
    """
    where = "" if can_all else "WHERE e.owner_id = ?"
    params = [] if can_all else [payload['username']]
    cursor.execute(base_sql + " " + where, params)
    enterprises = [dict(r) for r in cursor.fetchall()]

    today = datetime.now().date()
    results = []

    for ent in enterprises:
        eid = ent['id']
        # 关联客户 ID
        cursor.execute("""
            SELECT id FROM customers WHERE company LIKE ? OR company = ?
        """, (f'%{ent["name"]}%', ent['name']))
        cust_ids = [r['id'] for r in cursor.fetchall()]

        # 商机数据（未成交）
        business_amt = 0
        business_count = 0
        if cust_ids:
            placeholders = ','.join(['?'] * len(cust_ids))
            cursor.execute(f"""
                SELECT COALESCE(SUM(amount), 0) as total, COUNT(*) as cnt
                FROM business WHERE cust_id IN ({placeholders}) AND status NOT IN ('已成交', '已作废')
            """, cust_ids)
            br = cursor.fetchone()
            business_amt = float(br['total'] or 0)
            business_count = int(br['cnt'] or 0)

        # 合同数据
        contract_amt = 0
        if cust_ids:
            placeholders = ','.join(['?'] * len(cust_ids))
            cursor.execute(f"""
                SELECT COALESCE(SUM(total_amt), 0) as total FROM contracts WHERE cust_id IN ({placeholders})
            """, cust_ids)
            contract_amt = float(cursor.fetchone()['total'] or 0)

        # 上次拜访距今天数
        last_visit = ent.get('last_visit_date')
        days_since = 999
        if last_visit:
            try:
                days_since = (today - datetime.strptime(last_visit[:10], '%Y-%m-%d').date()).days
            except Exception:
                days_since = 999

        visit_count = int(ent.get('visit_count') or 0)
        status = ent.get('relationship_status') or '未接触'

        # === 评分计算 ===
        # 1. 商机价值（35分）：金额越高分越高，封顶35
        business_score = min(35, (business_amt / 1000000) * 5 + business_count * 2)

        # 2. 合同价值（20分）：已有合同金额越高分越高
        contract_score = min(20, (contract_amt / 1000000) * 3)

        # 3. 拜访间隔（25分）：30天内0分，30-90天线性，90天以上25分
        if days_since <= 30:
            interval_score = 0
        elif days_since >= 90:
            interval_score = 25
        else:
            interval_score = (days_since - 30) / 60 * 25

        # 4. 关系状态（10分）
        status_scores = {'合作中': 10, '已接触': 6, '意向中': 8, '未接触': 0, '流失': 3}
        status_score = status_scores.get(status, 0)

        # 5. 拜访频率（10分）：0次=5分（新客户需开拓），1-3次=10分，>3次递减
        if visit_count == 0:
            freq_score = 5
        elif visit_count <= 3:
            freq_score = 10
        elif visit_count <= 6:
            freq_score = 7
        else:
            freq_score = 4

        total_score = business_score + contract_score + interval_score + status_score + freq_score

        # 推荐理由
        reasons = []
        if business_amt > 0:
            reasons.append(f'待跟进商机 {business_count} 个（{business_amt/10000:.1f}万）')
        if contract_amt > 0:
            reasons.append(f'历史合同额 {contract_amt/10000:.1f}万，客户价值高')
        if days_since >= 60:
            reasons.append(f'已 {days_since} 天未拜访，需维护客户关系')
        if visit_count == 0:
            reasons.append('尚未拜访，建议开拓')
        if not reasons:
            reasons.append('常规客户关系维护')

        results.append({
            'id': eid,
            'name': ent['name'],
            'relationship_status': status,
            'contact_person': ent.get('contact_person', ''),
            'contact_info': ent.get('contact_info', ''),
            'location': ent.get('location', ''),
            'owner_name': ent.get('owner_name', ''),
            'last_visit_date': last_visit,
            'days_since_last_visit': days_since if days_since < 999 else None,
            'visit_count': visit_count,
            'business_amount': business_amt,
            'contract_amount': contract_amt,
            'score': round(total_score, 1),
            'reasons': reasons,
        })

    # 按分数降序
    results.sort(key=lambda x: x['score'], reverse=True)

    record_operation_log(
        payload['username'], '查询', '企业信息库',
        f'获取拜访推荐列表（共{len(results)}条）'
    )
    return jsonify({'code': 200, 'message': 'success', 'data': results})
